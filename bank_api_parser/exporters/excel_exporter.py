import logging
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side,
)
from openpyxl.utils import get_column_letter

from parsers.base_parser import APIMethod, ParseSnapshot

logger = logging.getLogger("exporter.excel")

OUTPUT_DIR = Path(__file__).parent.parent / "output"

BANK_DISPLAY_NAMES = {
    "tbank": "Т-Банк",
    "tochka": "Точка Банк",
    "alfabank": "Альфа-Банк",
    "sber": "Сбер",
}

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
ALT_ROW_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Calibri", size=10)
THIN_BORDER_SIDE = Side(border_style="thin", color="BFBFBF")
THIN_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=THIN_BORDER_SIDE,
)

METHOD_COLORS = {
    "GET": "70AD47",
    "POST": "4472C4",
    "PUT": "ED7D31",
    "PATCH": "FFC000",
    "DELETE": "FF0000",
}


def _style_header_row(ws, row_num: int, col_count: int):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def _style_data_row(ws, row_num: int, col_count: int, alt: bool = False):
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_num, column=col)
        if alt:
            cell.fill = ALT_ROW_FILL
        cell.font = BODY_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = THIN_BORDER


def _set_column_widths(ws, widths: list):
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


class ExcelExporter:
    def __init__(self, output_dir: Path = OUTPUT_DIR):
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export(
        self,
        snapshots: dict,
        filename: Optional[str] = None,
        diffs: Optional[list] = None,
    ) -> Path:
        """
        snapshots: {bank_name: ParseSnapshot or None}
        diffs: list of DiffResult objects (optional)
        Returns path to the generated Excel file.
        """
        if not filename:
            date_str = datetime.utcnow().strftime("%Y-%m-%d_%H-%M")
            filename = f"bank_api_report_{date_str}.xlsx"

        output_path = self.output_dir / filename
        wb = Workbook()

        # Remove default sheet
        default_sheet = wb.active
        wb.remove(default_sheet)

        # Sheet 1: Summary
        self._build_summary_sheet(wb, snapshots)

        # Sheet 2: Changes (if diffs provided)
        if diffs:
            self._build_diff_sheet(wb, diffs)

        # Sheets 3+: Details per bank
        bank_order = ["alfabank", "sber", "tbank", "tochka"]
        for bank in bank_order:
            snapshot = snapshots.get(bank)
            if snapshot is not None:
                self._build_detail_sheet(wb, snapshot)

        # Also add any banks not in the standard order
        for bank, snapshot in snapshots.items():
            if bank not in bank_order and snapshot is not None:
                self._build_detail_sheet(wb, snapshot)

        wb.save(output_path)
        logger.info("Excel report saved: %s", output_path)
        return output_path

    def _build_summary_sheet(self, wb: Workbook, snapshots: dict):
        ws = wb.create_sheet("Сводная таблица", 0)
        ws.freeze_panes = "A2"

        headers = ["Банк", "Кол-во сервисов", "Кол-во методов", "Дата парсинга", "Статус"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        _style_header_row(ws, 1, len(headers))
        ws.row_dimensions[1].height = 30

        bank_order = ["alfabank", "sber", "tbank", "tochka"]
        all_banks = bank_order + [b for b in snapshots if b not in bank_order]

        for row_idx, bank in enumerate(all_banks, 2):
            snapshot = snapshots.get(bank)
            display_name = BANK_DISPLAY_NAMES.get(bank, bank.title())
            alt = (row_idx % 2 == 0)

            if snapshot:
                parsed_date = snapshot.parsed_at[:10] if snapshot.parsed_at else ""
                ws.cell(row=row_idx, column=1, value=display_name)
                ws.cell(row=row_idx, column=2, value=snapshot.total_services)
                ws.cell(row=row_idx, column=3, value=snapshot.total_methods)
                ws.cell(row=row_idx, column=4, value=parsed_date)
                ws.cell(row=row_idx, column=5, value="OK")
                ws.cell(row=row_idx, column=5).font = Font(name="Calibri", size=10, color="375623", bold=True)
            else:
                ws.cell(row=row_idx, column=1, value=display_name)
                ws.cell(row=row_idx, column=5, value="ОШИБКА")
                ws.cell(row=row_idx, column=5).font = Font(name="Calibri", size=10, color="C00000", bold=True)

            _style_data_row(ws, row_idx, len(headers), alt)
            ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal="center", vertical="top")
            ws.cell(row=row_idx, column=3).alignment = Alignment(horizontal="center", vertical="top")
            ws.cell(row=row_idx, column=4).alignment = Alignment(horizontal="center", vertical="top")
            ws.cell(row=row_idx, column=5).alignment = Alignment(horizontal="center", vertical="top")

        _set_column_widths(ws, [25, 18, 16, 16, 10])
        ws.sheet_view.showGridLines = True

    def _build_detail_sheet(self, wb: Workbook, snapshot: ParseSnapshot):
        display_name = BANK_DISPLAY_NAMES.get(snapshot.bank, snapshot.bank.title())
        sheet_title = display_name[:31]  # Excel limit

        ws = wb.create_sheet(sheet_title)
        ws.freeze_panes = "A2"

        headers = ["Сервис", "HTTP-метод", "Путь (endpoint)", "Название метода", "Описание", "Ссылка"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)
        _style_header_row(ws, 1, len(headers))
        ws.row_dimensions[1].height = 30

        row_idx = 2
        for service_name in sorted(snapshot.services.keys()):
            methods = snapshot.services[service_name]
            for method in methods:
                if isinstance(method, dict):
                    from parsers.base_parser import APIMethod as AM
                    method = AM.from_dict(method)

                alt = (row_idx % 2 == 0)
                ws.cell(row=row_idx, column=1, value=service_name)
                ws.cell(row=row_idx, column=2, value=method.http_method)
                ws.cell(row=row_idx, column=3, value=method.path)
                ws.cell(row=row_idx, column=4, value=method.summary)
                ws.cell(row=row_idx, column=5, value=method.description[:500] if method.description else "")
                ws.cell(row=row_idx, column=6, value=method.url_on_portal)

                _style_data_row(ws, row_idx, len(headers), alt)

                # Color-code HTTP method
                http_color = METHOD_COLORS.get(method.http_method.upper(), "595959")
                ws.cell(row=row_idx, column=2).font = Font(
                    name="Calibri", size=10, bold=True, color=http_color
                )
                ws.cell(row=row_idx, column=2).alignment = Alignment(horizontal="center", vertical="top")

                # Hyperlink for portal URL
                if method.url_on_portal and method.url_on_portal.startswith("http"):
                    ws.cell(row=row_idx, column=6).hyperlink = method.url_on_portal
                    ws.cell(row=row_idx, column=6).font = Font(
                        name="Calibri", size=10, color="0563C1", underline="single"
                    )

                row_idx += 1

        _set_column_widths(ws, [30, 12, 45, 40, 60, 50])

        # Add auto-filter
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row_idx - 1}"

    # ── Diff sheet ────────────────────────────────────────────────────────────

    def _build_diff_sheet(self, wb: Workbook, diffs: list):
        """Build 'Изменения' sheet from a list of DiffResult objects."""
        ws = wb.create_sheet("Изменения", 1)
        ws.freeze_panes = "A2"

        ADDED_FILL   = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        REMOVED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        CHANGED_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        SECTION_FILL = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
        SECTION_FONT = Font(name="Calibri", bold=True, size=11)

        headers = ["Тип изменения", "Банк", "Сервис", "HTTP-метод", "Путь (endpoint)", "Описание", "Период"]
        for col, h in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=h)
        _style_header_row(ws, 1, len(headers))
        ws.row_dimensions[1].height = 28

        row = 2
        has_changes = False

        for diff in diffs:
            if not diff.has_changes:
                continue
            has_changes = True

            bank_name = BANK_DISPLAY_NAMES.get(diff.bank, diff.bank.title())
            period = f"{diff.date_old} → {diff.date_new}"

            def _write_change(change_type: str, service: str, verb: str, path: str, desc: str, fill):
                nonlocal row
                ws.cell(row=row, column=1, value=change_type)
                ws.cell(row=row, column=2, value=bank_name)
                ws.cell(row=row, column=3, value=service)
                ws.cell(row=row, column=4, value=verb)
                ws.cell(row=row, column=5, value=path)
                ws.cell(row=row, column=6, value=desc[:300] if desc else "")
                ws.cell(row=row, column=7, value=period)
                for col in range(1, len(headers) + 1):
                    cell = ws.cell(row=row, column=col)
                    cell.fill = fill
                    cell.font = BODY_FONT
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cell.border = THIN_BORDER
                ws.cell(row=row, column=4).alignment = Alignment(horizontal="center", vertical="top")
                row += 1

            # New services
            for svc in diff.added_services:
                _write_change("+ Новый сервис", svc, "", "", "", ADDED_FILL)

            # Removed services
            for svc in diff.removed_services:
                _write_change("− Удалён сервис", svc, "", "", "", REMOVED_FILL)

            # Added methods
            for m in diff.added_methods:
                _write_change("+ Новый метод", m.service, m.http_method, m.path, m.summary, ADDED_FILL)

            # Removed methods
            for m in diff.removed_methods:
                _write_change("− Удалён метод", m.service, m.http_method, m.path, m.summary, REMOVED_FILL)

            # Changed fields
            for m in diff.changed_methods:
                changes_str = ", ".join(m.changed_fields[:10])
                _write_change(
                    "~ Изменены поля 200 OK", m.service, m.http_method, m.path,
                    f"Изменения: {changes_str}", CHANGED_FILL,
                )

        if not has_changes:
            ws.cell(row=2, column=1, value="Изменений не обнаружено")
            ws.cell(row=2, column=1).font = Font(name="Calibri", size=11, italic=True, color="595959")
            row = 3

        _set_column_widths(ws, [22, 16, 35, 12, 45, 55, 22])
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{row - 1}"

        # Summary block above the data
        ws.insert_rows(1)
        ws.insert_rows(1)
        ws.cell(row=1, column=1, value="ОТЧЁТ ОБ ИЗМЕНЕНИЯХ API")
        ws.cell(row=1, column=1).font = Font(name="Calibri", bold=True, size=13)
        ws.cell(row=2, column=1, value=f"Сформирован: {datetime.utcnow().strftime('%d.%m.%Y %H:%M')} UTC")
        ws.cell(row=2, column=1).font = Font(name="Calibri", size=10, italic=True, color="595959")
        ws.row_dimensions[1].height = 22
        ws.row_dimensions[2].height = 16
