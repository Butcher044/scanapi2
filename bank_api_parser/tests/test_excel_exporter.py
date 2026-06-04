"""Unit tests for ExcelExporter."""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from openpyxl import load_workbook

from exporters.excel_exporter import ExcelExporter
from parsers.base_parser import APIMethod, ParseSnapshot


def _make_snapshot(bank: str, num_services: int = 3, methods_per_service: int = 4) -> ParseSnapshot:
    services = {}
    total = 0
    for s_idx in range(num_services):
        svc = f"Service {s_idx + 1}"
        methods = []
        for m_idx in range(methods_per_service):
            verbs = ["GET", "POST", "PUT", "DELETE"]
            verb = verbs[m_idx % len(verbs)]
            methods.append(APIMethod(
                bank=bank,
                service_name=svc,
                http_method=verb,
                path=f"/api/v1/{svc.lower().replace(' ', '_')}/{m_idx}",
                summary=f"Method {m_idx + 1}",
                description=f"Description for method {m_idx + 1} in {svc}",
                response_200_fields=[f"field_{i}" for i in range(5)],
                parsed_at="2026-05-28T10:00:00",
                url_on_portal=f"https://example.com/{bank}/{svc}#{m_idx}",
            ))
        services[svc] = methods
        total += len(methods)

    return ParseSnapshot(
        bank=bank,
        parsed_at="2026-05-28T10:00:00",
        services=services,
        total_services=num_services,
        total_methods=total,
    )


def test_export_creates_file():
    with tempfile.TemporaryDirectory() as tmpdir:
        exporter = ExcelExporter(output_dir=Path(tmpdir))
        snapshots = {"tbank": _make_snapshot("tbank")}
        path = exporter.export(snapshots, filename="test.xlsx")
        assert path.exists()
        assert path.suffix == ".xlsx"


def test_summary_sheet_exists():
    with tempfile.TemporaryDirectory() as tmpdir:
        exporter = ExcelExporter(output_dir=Path(tmpdir))
        snapshots = {
            "tbank": _make_snapshot("tbank", 2, 3),
            "tochka": _make_snapshot("tochka", 2, 3),
        }
        path = exporter.export(snapshots, filename="test.xlsx")
        wb = load_workbook(path)
        assert "Сводная таблица" in wb.sheetnames


def test_summary_has_correct_row_count():
    with tempfile.TemporaryDirectory() as tmpdir:
        exporter = ExcelExporter(output_dir=Path(tmpdir))
        banks = ["tbank", "tochka", "alfabank"]
        snapshots = {b: _make_snapshot(b, 2, 2) for b in banks}
        path = exporter.export(snapshots, filename="test.xlsx")
        wb = load_workbook(path)
        ws = wb["Сводная таблица"]
        # Header row + 3 bank rows
        assert ws.max_row >= 4


def test_detail_sheets_created():
    with tempfile.TemporaryDirectory() as tmpdir:
        exporter = ExcelExporter(output_dir=Path(tmpdir))
        snapshots = {
            "tbank": _make_snapshot("tbank", 2, 3),
            "tochka": _make_snapshot("tochka", 2, 3),
            "alfabank": _make_snapshot("alfabank", 2, 3),
            "sber": _make_snapshot("sber", 2, 3),
        }
        path = exporter.export(snapshots, filename="test.xlsx")
        wb = load_workbook(path)
        assert len(wb.sheetnames) == 5  # 1 summary + 4 detail


def test_detail_sheet_method_count():
    with tempfile.TemporaryDirectory() as tmpdir:
        exporter = ExcelExporter(output_dir=Path(tmpdir))
        snap = _make_snapshot("tbank", num_services=3, methods_per_service=4)  # 12 methods total
        path = exporter.export({"tbank": snap}, filename="test.xlsx")
        wb = load_workbook(path)
        # Find T-Bank detail sheet
        detail_sheet = None
        for name in wb.sheetnames:
            if "Банк" in name or "tbank" in name.lower() or "Т-Банк" in name:
                detail_sheet = wb[name]
                break
        assert detail_sheet is not None
        # Header + 12 rows
        assert detail_sheet.max_row == 13


def test_none_snapshot_shown_in_summary():
    with tempfile.TemporaryDirectory() as tmpdir:
        exporter = ExcelExporter(output_dir=Path(tmpdir))
        snapshots = {
            "tbank": _make_snapshot("tbank"),
            "sber": None,
        }
        path = exporter.export(snapshots, filename="test.xlsx")
        wb = load_workbook(path)
        ws = wb["Сводная таблица"]
        # Check that Сбер row has "ОШИБКА" status
        found_error = False
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and "Сбер" in str(row[0]):
                if row[4] and "ОШИБКА" in str(row[4]):
                    found_error = True
        assert found_error


def test_empty_snapshots():
    with tempfile.TemporaryDirectory() as tmpdir:
        exporter = ExcelExporter(output_dir=Path(tmpdir))
        path = exporter.export({}, filename="test.xlsx")
        wb = load_workbook(path)
        assert "Сводная таблица" in wb.sheetnames


if __name__ == "__main__":
    tests = [v for k, v in list(globals().items()) if k.startswith("test_")]
    passed = 0
    failed = 0
    for test_fn in tests:
        try:
            test_fn()
            print(f"  PASS  {test_fn.__name__}")
            passed += 1
        except Exception as exc:
            import traceback
            print(f"  FAIL  {test_fn.__name__}: {exc}")
            traceback.print_exc()
            failed += 1
    print(f"\n{passed} passed, {failed} failed")
    sys.exit(0 if failed == 0 else 1)
